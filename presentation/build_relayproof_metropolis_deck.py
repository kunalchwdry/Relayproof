#!/usr/bin/env python3
"""Build a 7-slide RelayProof / Monad Metropolis problem-fit deck."""
from __future__ import annotations

from pathlib import Path
import math

import cairosvg
import numpy as np
from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "presentation_assets" / "monad"
OUT = ROOT / "presentation"
OUT.mkdir(parents=True, exist_ok=True)

PPTX_PATH = OUT / "relayproof_monad_metropolis_pitch_deck.pptx"
BG_PATH = OUT / "relayproof_video_theme_bg.png"
MONAD_ICON = ASSETS / "monad-official-avatar.png"
METROPOLIS_SVG = ASSETS / "metropolis-wordmark.svg"
METROPOLIS_PNG = ASSETS / "metropolis-wordmark.png"

SW, SH = 13.333, 7.5
# RelayProof video theme, with a restrained Monad purple brand accent.
C = {
    "bg": "0B201B", "bg2": "10291F", "grid": "16342A", "card": "15362B",
    "card2": "1B4032", "card3": "203E34", "line": "345644", "paper": "F3F1E8",
    "muted": "BAC4B7", "dim": "8FA294", "gold": "CFAC74", "mint": "99C5A2",
    "blue": "88B7BA", "coral": "DD8B78", "lav": "C3ADD0", "purple": "6E54FF",
    "white": "FFFFFF", "ink": "0A1713",
}

FONT = "Arial"
MONO = "Consolas"


def rgb(name: str) -> RGBColor:
    h = C.get(name, name).replace("#", "")
    return RGBColor.from_string(h.upper())


def ensure_assets() -> None:
    if not MONAD_ICON.exists():
        raise FileNotFoundError(f"Official Monad logo asset missing: {MONAD_ICON}")
    if not METROPOLIS_PNG.exists():
        if not METROPOLIS_SVG.exists():
            raise FileNotFoundError(f"Official Metropolis wordmark missing: {METROPOLIS_SVG}")
        cairosvg.svg2png(url=str(METROPOLIS_SVG), write_to=str(METROPOLIS_PNG), output_width=1500)
    if not BG_PATH.exists():
        make_background(BG_PATH)


def make_background(path: Path) -> None:
    w, h = 1600, 900
    yy, xx = np.mgrid[0:h, 0:w]
    mix = (yy / (h - 1))[..., None]
    top = np.array([11, 32, 27], dtype=np.float32)
    bottom = np.array([7, 23, 19], dtype=np.float32)
    px = top * (1 - mix) + bottom * mix
    glow = np.exp(-(((xx - 1400) / 560) ** 2 + ((yy - 110) / 370) ** 2) * 1.7)[..., None]
    px += glow * np.array([28, 27, 11], dtype=np.float32)
    px = np.clip(px, 0, 255).astype(np.uint8)
    im = Image.fromarray(px, "RGB")
    d = ImageDraw.Draw(im)
    for x in range(25, w, 80): d.line((x, 0, x, h), fill=(18, 43, 34), width=1)
    for y in range(22, h, 80): d.line((0, y, w, y), fill=(18, 43, 34), width=1)
    im.save(path, optimize=True)


def add_bg(slide):
    slide.shapes.add_picture(str(BG_PATH), 0, 0, width=Inches(SW), height=Inches(SH))


def add_rect(slide, x, y, w, h, fill="card", line="line", radius=True, line_width=0.8):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                  Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line); shape.line.width = Pt(line_width)
    return shape


def add_line(slide, x1, y1, x2, y2, color="line", width=1.0):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = rgb(color); shape.line.width = Pt(width)
    return shape


def text_box(slide, x, y, w, h, text, size=16, color="paper", bold=False,
             font=FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
             margin=0.02, italic=False, tracking=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(margin); tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin); tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.alignment = align
        p.space_before = Pt(0); p.space_after = Pt(1)
        p.font.name = font; p.font.size = Pt(size); p.font.bold = bold; p.font.italic = italic
        p.font.color.rgb = rgb(color)
    return shape


def pill(slide, x, y, label, color="gold", width=None, size=9.2, fill="card2"):
    w = width or max(0.92, len(label) * size * 0.0077 + 0.34)
    h = 0.29
    add_rect(slide, x, y, w, h, fill=fill, line=color, radius=True, line_width=0.65)
    text_box(slide, x+0.08, y+0.035, w-0.16, h-0.055, label, size=size, color=color,
             font=MONO, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
    return w


def add_icon(slide, x, y, size=0.34):
    slide.shapes.add_picture(str(MONAD_ICON), Inches(x), Inches(y), width=Inches(size), height=Inches(size))


def add_wordmark(slide, x, y, w):
    with Image.open(METROPOLIS_PNG) as im:
        h = w * im.height / im.width
    slide.shapes.add_picture(str(METROPOLIS_PNG), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    return h


def header(slide, section: str, page: int):
    text_box(slide, 0.62, 0.24, 2.85, 0.22, "RELAYPROOF  /  PRODUCT BRIEF", size=10.2, color="muted", bold=True, font=MONO, margin=0)
    add_icon(slide, 9.96, 0.19, 0.35)
    text_box(slide, 10.40, 0.255, 2.25, 0.18, "MONAD · METROPOLIS", size=8.6, color="muted", bold=True, font=MONO, align=PP_ALIGN.RIGHT, margin=0)
    add_line(slide, 0.62, 0.60, 12.70, 0.60, color="line", width=0.7)
    add_line(slide, 0.62, 7.12, 12.70, 7.12, color="line", width=0.7)
    text_box(slide, 0.64, 7.19, 8.8, 0.16, f"{section.upper()}  ·  TRACK-FIT CONCEPT / PROTOTYPE STATUS", size=8.2, color="dim", font=MONO, margin=0)
    text_box(slide, 12.0, 7.18, 0.63, 0.18, f"{page:02d} / 07", size=8.2, color="dim", font=MONO, align=PP_ALIGN.RIGHT, margin=0)


def slide_title(slide, kicker: str, title: str, subtitle: str | None = None):
    text_box(slide, 0.68, 0.78, 4.6, 0.23, kicker.upper(), size=9.5, color="gold", font=MONO, bold=True, margin=0)
    text_box(slide, 0.68, 1.04, 12.0, 0.55, title, size=25.5, color="paper", bold=True, margin=0)
    if subtitle:
        text_box(slide, 0.70, 1.59, 11.9, 0.38, subtitle, size=12.2, color="muted", margin=0)


def card(slide, x, y, w, h, label, title, body, accent="mint", label_width=None,
         body_size=12.0, title_size=16.2, fill="card"):
    add_rect(slide, x, y, w, h, fill=fill, line="line", radius=True, line_width=0.7)
    # slim accent at the left edge
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x+0.02), Inches(y+0.18), Inches(0.04), Inches(h-0.36))
    stripe.fill.solid(); stripe.fill.fore_color.rgb = rgb(accent); stripe.line.fill.background()
    pill(slide, x+0.19, y+0.15, label.upper(), color=accent, width=label_width, size=8.1)
    text_box(slide, x+0.19, y+0.54, w-0.37, 0.42, title, size=title_size, color="paper", bold=True, margin=0)
    text_box(slide, x+0.19, y+1.02, w-0.38, h-1.14, body, size=body_size, color="muted", margin=0.01)


def bullet(slide, x, y, w, text, color="paper", dot="mint", size=12.2, h=0.42, bold=False):
    dotshape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y+0.105), Inches(0.085), Inches(0.085))
    dotshape.fill.solid(); dotshape.fill.fore_color.rgb = rgb(dot); dotshape.line.fill.background()
    return text_box(slide, x+0.19, y, w-0.19, h, text, size=size, color=color, bold=bold, margin=0.0)


def source_link(slide, x, y, label, url, width=5.8):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(0.22))
    tf = shape.text_frame; tf.clear(); tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = label
    r.font.name = FONT; r.font.size = Pt(8.5); r.font.color.rgb = rgb("blue")
    r.hyperlink.address = url
    return shape


def make_slide(prs: Presentation, section: str, page: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide); header(slide, section, page)
    return slide


def slide1(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s)
    text_box(s, 0.70, 0.40, 3.7, 0.23, "PRODUCT + PROBLEM-FIT DECK", size=10, color="muted", bold=True, font=MONO, margin=0)
    add_icon(s, 9.05, 0.35, 0.48)
    text_box(s, 9.62, 0.43, 0.98, 0.20, "MONAD", size=10, color="paper", bold=True, font=MONO, margin=0)
    add_wordmark(s, 10.75, 0.43, 1.88)
    add_line(s, 0.68, 0.88, 12.68, 0.88, color="line", width=0.8)

    pill(s, 0.78, 1.38, "METROPOLIS · TRACK 04", color="gold", width=2.34, size=9.2)
    text_box(s, 0.78, 1.92, 5.75, 0.84, "RelayProof", size=40, color="paper", bold=True, margin=0)
    text_box(s, 0.80, 2.74, 5.55, 1.03, "A clearer evidence trail for audio & video.", size=27.5, color="mint", bold=True, margin=0)
    text_box(s, 0.82, 3.92, 5.38, 0.72, "Signed source claims · exact-file checks · local content review · Chrome extension", size=16.1, color="muted", margin=0)
    add_rect(s, 0.80, 5.02, 5.55, 0.98, fill="card", line="gold", radius=True, line_width=0.9)
    pill(s, 1.00, 5.19, "CURRENT BOUNDARY", color="gold", width=1.55, size=8.0)
    text_box(s, 1.00, 5.53, 5.02, 0.36, "Exact-byte integrity today — not re-encoding resilience.", size=10.7, color="paper", bold=True, margin=0)
    text_box(s, 0.82, 6.32, 5.9, 0.40, "Track-fit concept only · no Monad endorsement, entrant status, or onchain deployment claimed.", size=9.2, color="dim", font=MONO, margin=0)

    add_rect(s, 6.95, 1.34, 5.68, 4.91, fill="card", line="line", radius=True, line_width=0.8)
    pill(s, 7.24, 1.62, "THE EVIDENCE PATH", color="mint", width=1.62, size=8.4)
    text_box(s, 7.24, 2.04, 5.05, 0.43, "One file. One signed claim. Clear limits.", size=15.4, color="paper", bold=True, margin=0)
    # Three native diagram nodes
    nodes = [(7.35, "01", "PUBLISHER", "selects media\n+ states origin", "blue"),
             (9.20, "02", "SIDECAR", "SHA-256 +\nsigned fields", "gold"),
             (11.05, "03", "RECIPIENT", "checks bytes\n+ reviews", "mint")]
    for x, num, name, body, accent in nodes:
        add_rect(s, x, 2.88, 1.43, 1.38, fill="card2", line=accent, radius=True, line_width=0.9)
        pill(s, x+0.12, 3.02, num, color=accent, width=0.42, size=8.2)
        text_box(s, x+0.12, 3.42, 1.18, 0.25, name, size=9.5, color="paper", bold=True, font=MONO, margin=0)
        text_box(s, x+0.12, 3.76, 1.20, 0.43, body, size=9.7, color="muted", margin=0)
    # Connect nodes
    for x1, x2 in [(8.82, 9.13), (10.67, 10.98)]:
        add_line(s, x1, 3.56, x2, 3.56, color="gold", width=2.0)
        tri = s.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, Inches(x2-0.035), Inches(3.49), Inches(0.14), Inches(0.14))
        tri.rotation = 45; tri.fill.solid(); tri.fill.fore_color.rgb = rgb("gold"); tri.line.fill.background()
    add_line(s, 7.28, 4.62, 12.32, 4.62, color="line", width=0.8)
    text_box(s, 7.27, 4.85, 2.25, 0.29, "CRYPTOGRAPHIC CHECK", size=9.1, color="mint", font=MONO, bold=True, margin=0)
    text_box(s, 7.27, 5.18, 4.95, 0.62, "A valid signature and matching digest answer narrow questions; they do not prove real-world identity or truth.", size=11.3, color="muted", margin=0)
    add_line(s, 0.68, 7.12, 12.68, 7.12, color="line", width=0.8)
    text_box(s, 0.70, 7.20, 9.8, 0.16, "RELAYPROOF  ·  EVIDENCE FIRST; VERDICTS NEVER WITHOUT PROOF", size=8.1, color="dim", font=MONO, margin=0)
    text_box(s, 12.00, 7.20, 0.63, 0.16, "01 / 07", size=8.1, color="dim", font=MONO, align=PP_ALIGN.RIGHT, margin=0)
    return s


def slide2(prs):
    s = make_slide(prs, "Official brief", 2)
    slide_title(s, "01 / METROPOLIS", "The brief: four tracks, one selected problem", "Track 04: Trust, Identity & AI Infrastructure · Global online, 1 Sep–13 Oct 2026.")
    add_rect(s, 0.72, 1.98, 11.92, 0.77, fill="card2", line="gold", radius=True, line_width=0.8)
    text_box(s, 0.98, 2.08, 11.35, 0.57, "“AI has made identity, provenance, data ownership, and agent trust urgent infrastructure problems.”", size=14.6, color="paper", bold=True, margin=0)
    track_data = [
        ("01", "ONCHAIN FINANCE & TRADING", "Onchain markets, lending, trading", "blue"),
        ("02", "CONSUMER PRODUCTS & PAYMENTS", "Invisible, programmable payments", "gold"),
        ("03", "SOCIAL, ATTENTION & CULTURE", "Social graphs, incentives, culture", "lav"),
        ("04", "TRUST, IDENTITY & AI INFRASTRUCTURE", "Identity, provenance, agents", "mint"),
    ]
    x0, gap, cw = 0.72, 0.15, 2.88
    for i, (num, name, descr, accent) in enumerate(track_data):
        x = x0 + i*(cw+gap)
        add_rect(s, x, 2.93, cw, 0.92, fill="card2" if num == "04" else "card", line=accent, radius=True, line_width=1.05 if num == "04" else 0.65)
        pill(s, x+0.15, 3.07, f"TRACK {num}", color=accent, width=0.90, size=7.5)
        text_box(s, x+0.15, 3.43, cw-0.30, 0.20, ("TRUST & AI INFRASTRUCTURE" if num == "04" else name), size=8.4, color="paper", bold=True, font=MONO, margin=0)
        text_box(s, x+0.15, 3.65, cw-0.30, 0.15, descr, size=8.9, color="muted", margin=0)
    # Track 04 detailed prompt and fit.
    add_rect(s, 0.72, 4.07, 6.00, 2.55, fill="card", line="line", radius=True, line_width=0.7)
    pill(s, 0.95, 4.28, "TRACK 04 · EXAMPLE PROMPTS", color="mint", width=2.20, size=8.0)
    bullet(s, 0.98, 4.73, 5.42, "Passkey-native P-256 + WebAuthn accounts, without a seed phrase.", dot="blue", size=11.2, h=0.41)
    bullet(s, 0.98, 5.25, 5.42, "Agent identity and reputation under ERC-8004.", dot="gold", size=11.2, h=0.38)
    bullet(s, 0.98, 5.73, 5.42, "Generated-media provenance that survives re-encoding.", dot="coral", size=11.2, h=0.42)
    add_rect(s, 6.92, 4.07, 5.72, 2.55, fill="card", line="gold", radius=True, line_width=0.8)
    pill(s, 7.17, 4.28, "RELAYPROOF FIT", color="gold", width=1.35, size=8.0)
    text_box(s, 7.18, 4.70, 5.10, 0.60, "Targets the media-provenance prompt — but does not yet satisfy its re-encoding requirement.", size=12.4, color="paper", bold=True, margin=0)
    text_box(s, 7.18, 5.38, 5.02, 0.98, "Today: self-asserted publisher claim + signed SHA-256 of exact bytes. Re-encoding changes the digest; Chromaprint is only an uncalibrated audio-similarity hint. No C2PA/SBR recovery or Monad anchor is connected.", size=10.7, color="muted", margin=0)
    source_link(s, 0.76, 6.75, "Official Metropolis prompt: monad.xyz/developers/hackathons/metropolis", "https://monad.xyz/developers/hackathons/metropolis", width=7.8)
    return s


def slide3(prs):
    s = make_slide(prs, "Problem", 3)
    slide_title(s, "02 / USER PROBLEM", "Media loses context as it moves", "A forwarded voice note or clip can arrive without dependable, portable source evidence.")
    # Three-stage media journey
    flow = [(0.78, "SOURCE", "Publisher selects\naudio / video", "blue"),
            (4.64, "TRANSFORM", "Transcode · trim\nrepost · screenshot", "gold"),
            (8.50, "RECIPIENT", "Needs to assess\nsource + media", "mint")]
    for x, tag, body, accent in flow:
        add_rect(s, x, 2.06, 3.24, 1.02, fill="card", line=accent, radius=True, line_width=0.8)
        pill(s, x+0.18, 2.22, tag, color=accent, width=1.05, size=8.1)
        text_box(s, x+0.18, 2.58, 2.88, 0.45, body, size=11.8, color="paper", bold=True, margin=0)
    for x1, x2 in [(4.05,4.54),(7.91,8.40)]:
        add_line(s, x1, 2.57, x2, 2.57, color="gold", width=1.7)
        tri = s.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, Inches(x2-0.02), Inches(2.50), Inches(0.13), Inches(0.13))
        tri.rotation=45; tri.fill.solid(); tri.fill.fore_color.rgb=rgb("gold"); tri.line.fill.background()
    # Core problems
    card(s, 0.78, 3.42, 3.78, 1.70, "INTEGRITY", "Bytes change", "Exact SHA-256 is strong for the original file; a transcode normally produces a different digest.", accent="gold", body_size=10.0, title_size=14.2)
    card(s, 4.78, 3.42, 3.78, 1.70, "IDENTITY", "A key is not a person", "The sidecar embeds a self-asserted public key. A valid signature does not establish real-world publisher identity.", accent="blue", body_size=10.0, title_size=14.2)
    card(s, 8.78, 3.42, 3.78, 1.70, "CONTENT", "Evidence is not truth", "Transcript, OCR, and phrase flags can miss context; they do not fact-check claims or approve SEBI compliance.", accent="coral", body_size=10.0, title_size=14.2)
    add_rect(s, 0.78, 5.28, 11.78, 1.08, fill="card2", line="mint", radius=True, line_width=0.8)
    pill(s, 1.02, 5.48, "PRODUCT HYPOTHESIS", color="mint", width=1.78, size=8.0)
    text_box(s, 1.02, 5.82, 11.05, 0.47, "Separate source claims, exact-byte integrity, candidate similarity, and human review—showing what each signal can and cannot prove.", size=12.6, color="paper", bold=True, margin=0)
    text_box(s, 0.82, 6.61, 11.4, 0.21, "Critical gap: post-re-encoding provenance is the official Track 04 ambition; the prototype has not delivered that capability.", size=10.4, color="coral", bold=True, font=MONO, margin=0)
    return s


def slide4(prs):
    s = make_slide(prs, "Product", 4)
    slide_title(s, "03 / PRODUCT", "Six features, with clear evidence boundaries", "The prototype keeps cryptographic checks separate from optional matching and human content review.")
    card(s, 0.72, 2.05, 3.83, 1.78, "ISSUE", "Sign a source record", "Audio/video up to 128 MB. Local ECDSA P-256 key signs a self-asserted claim and whole-file SHA-256; export JSON sidecar.", accent="mint", body_size=10.5, title_size=15.0)
    card(s, 4.75, 2.05, 3.83, 1.78, "VERIFY", "Two separate checks", "Verify the signature against its included public key, then compare exact-file bytes. A local key pin is only a reminder on this device.", accent="blue", body_size=10.5, title_size=15.0)
    card(s, 8.78, 2.05, 3.83, 1.78, "ARCHIVE", "Browser-scoped records", "Reopen or export issued sidecars from this browser. Browser storage is not a backup; keep sidecars separately.", accent="gold", body_size=10.5, title_size=15.0)
    card(s, 0.72, 4.02, 3.83, 1.92, "CHROMAPRINT", "Candidate audio match", "Optional self-hosted fingerprint of up to 120 seconds. Similarity is uncalibrated; no threshold, identity, truth, or authentication verdict.", accent="lav", body_size=10.4, title_size=14.7)
    card(s, 4.75, 4.02, 3.83, 1.92, "CONTENT REVIEW", "Transcript + preflight", "Import captions or text; optional local Whisper. Limited English phrase rules surface selected risk language, numbers, and dates for human checking.", accent="gold", body_size=10.4, title_size=14.7)
    card(s, 8.78, 4.02, 3.83, 1.92, "VIDEO OCR", "Sampled stills only", "Up to 24 spaced frames sent to local Tesseract after a click. Timestamped English OCR can miss brief, small, stylized, or non-English text.", accent="coral", body_size=10.4, title_size=14.7)
    add_rect(s, 0.74, 6.18, 11.88, 0.44, fill="card2", line="line", radius=True, line_width=0.55)
    text_box(s, 0.94, 6.29, 11.46, 0.20, "No feature returns “SEBI approved” or “good to go.” Risk rules vary by speaker role, content type, date, and context; human review is required.", size=10.1, color="paper", bold=True, margin=0)
    return s


def slide5(prs):
    s = make_slide(prs, "Architecture + extension", 5)
    slide_title(s, "04 / EXPERIENCE", "Local-first checks, user-initiated media paths", "Exact checks stay in the browser; optional media tasks go to the user’s own local service.")
    # Architecture row
    boxes = [(0.76, "PUBLISHER / USER", "Select media + sidecar\nChoose task explicitly", "blue"),
             (4.77, "BROWSER / EXTENSION", "Sidecar signature\nSHA-256 exact-file check", "mint"),
             (8.78, "LOCAL API · OPTIONAL", "Chromaprint · Whisper\nTesseract on sampled stills", "gold")]
    for x, label, body, accent in boxes:
        add_rect(s, x, 2.02, 3.46, 1.08, fill="card", line=accent, radius=True, line_width=0.8)
        pill(s, x+0.18, 2.18, label, color=accent, width=2.1 if len(label)>15 else 1.7, size=7.7)
        text_box(s, x+0.18, 2.54, 3.10, 0.49, body, size=10.7, color="paper", bold=True, margin=0)
    for a,b in [(4.28,4.69),(8.29,8.70)]:
        add_line(s,a,2.56,b,2.56,color="gold",width=1.7)
        tri=s.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, Inches(b-0.02), Inches(2.49), Inches(0.13), Inches(0.13)); tri.rotation=45; tri.fill.solid(); tri.fill.fore_color.rgb=rgb("gold"); tri.line.fill.background()
    # Browser local callout
    text_box(s, 4.95, 3.22, 3.10, 0.23, "HASH / SIGNATURE CHECKS DO NOT UPLOAD MEDIA", size=8.5, color="mint", font=MONO, align=PP_ALIGN.CENTER, margin=0)
    # Two extension use paths
    card(s, 0.76, 3.61, 5.72, 1.72, "PATH A · LOCAL FILE", "Exact check, then optional compare", "Select a media file + sidecar. Verify signature and exact hash locally (up to 128 MB); if bytes differ, request an optional candidate comparison when a signed fingerprint exists.", accent="blue", body_size=10.6, title_size=14.8)
    card(s, 6.73, 3.61, 5.51, 1.72, "PATH B · CURRENT TAB", "Capture 15 / 30 / 60 seconds", "Start playback, load a sidecar with a Chromaprint reference, click capture, and keep the popup open. Captures current tab audio; it does not download the source or exact-match its bytes.", accent="gold", body_size=10.6, title_size=14.8)
    add_rect(s, 0.76, 5.68, 11.48, 0.68, fill="card2", line="line", radius=True, line_width=0.65)
    text_box(s, 0.98, 5.77, 10.97, 0.50, "Privacy boundary: selected media is sent only to a self-hosted local backend for opted-in fingerprinting/transcription; OCR sends sampled stills. No AcoustID lookup or cloud inference is configured.", size=9.7, color="paper", margin=0)
    text_box(s, 0.82, 6.59, 11.3, 0.20, "Manifest V3 extension · Optional services need native Chromaprint / Tesseract and faster-whisper; runtime setup not yet validated.", size=9.0, color="dim", font=MONO, margin=0)
    return s


def slide6(prs):
    s = make_slide(prs, "Fit + roadmap", 6)
    slide_title(s, "05 / TRACK FIT", "What exists, what is missing, and how to close the gap", "The prompt’s hardest part is provenance that survives re-encoding—not merely publishing a hash on a chain.")
    # Status columns
    add_rect(s, 0.74, 2.03, 3.82, 2.10, fill="card", line="mint", radius=True, line_width=0.8)
    pill(s, 0.97, 2.22, "IMPLEMENTED IN PROTOTYPE", color="mint", width=2.28, size=7.8)
    bullet(s, 0.99, 2.67, 3.25, "ECDSA-signed sidecars + exact SHA-256.", dot="mint", size=10.5, h=0.34)
    bullet(s, 0.99, 3.09, 3.25, "Archive, verifier UI, local review flows.", dot="mint", size=10.5, h=0.34)
    bullet(s, 0.99, 3.51, 3.25, "MV3 extension code + build checks pass.", dot="mint", size=9.8, h=0.34)

    add_rect(s, 4.76, 2.03, 3.82, 2.10, fill="card", line="gold", radius=True, line_width=0.8)
    pill(s, 4.99, 2.22, "OPTIONAL · NOT RUNTIME-VALIDATED", color="gold", width=2.62, size=7.6)
    bullet(s, 5.01, 2.67, 3.25, "Chromaprint needs local native libraries.", dot="gold", size=10.5, h=0.34)
    bullet(s, 5.01, 3.09, 3.25, "Whisper / Tesseract need local services.", dot="gold", size=10.5, h=0.34)
    bullet(s, 5.01, 3.51, 3.25, "No Chrome runtime test yet; API offline.", dot="coral", size=10.5, h=0.34)

    add_rect(s, 8.78, 2.03, 3.82, 2.10, fill="card", line="coral", radius=True, line_width=0.8)
    pill(s, 9.01, 2.22, "NOT CONNECTED", color="coral", width=1.52, size=7.8)
    bullet(s, 9.03, 2.67, 3.25, "C2PA / watermark / SBR recovery.", dot="coral", size=10.5, h=0.34)
    bullet(s, 9.03, 3.09, 3.25, "Monad contract, registry, revocation.", dot="coral", size=10.5, h=0.34)
    bullet(s, 9.03, 3.51, 3.25, "No post-transcode provenance proof.", dot="coral", size=10.5, h=0.34)

    add_rect(s, 0.74, 4.38, 7.46, 1.85, fill="card2", line="line", radius=True, line_width=0.7)
    pill(s, 0.98, 4.57, "ROADMAP TO A TESTABLE TRACK-04 CLAIM", color="blue", width=2.86, size=7.7)
    bullet(s, 1.00, 4.98, 6.90, "1 · Integrate SBR after sanitized API docs + test vectors are provided.", dot="blue", size=9.8, h=0.34)
    bullet(s, 1.00, 5.39, 6.90, "2 · Test Opus/AAC, trims, noise, and edits; report misses and false matches.", dot="blue", size=9.8, h=0.34)
    bullet(s, 1.00, 5.80, 6.90, "3 · If useful, anchor digests/status on Monad; keep media/private data off-chain.", dot="blue", size=9.8, h=0.34)

    add_rect(s, 8.42, 4.38, 4.18, 1.85, fill="card", line="gold", radius=True, line_width=0.7)
    pill(s, 8.66, 4.57, "POSITIONING · HYPOTHESIS", color="gold", width=2.16, size=7.7)
    text_box(s, 8.68, 5.02, 3.62, 0.96, "Adjacent C2PA/SBR, watermark, browser verification, and Metropolis provenance work already exist. Do not claim category-first. Test whether explicit evidence boundaries + an audio/video recipient workflow add value.", size=10.3, color="muted", margin=0)
    text_box(s, 0.82, 6.51, 11.4, 0.21, "Scope choice: media provenance only; passkey and ERC-8004 examples are separate directions, not requirements for this build.", size=9.3, color="coral", bold=True, font=MONO, margin=0)
    return s


def slide7(prs):
    s = make_slide(prs, "Evidence + next steps", 7)
    slide_title(s, "06 / DEMO + VALIDATION", "Show the boundary—not just the happy path", "Make the current product easy to evaluate, then collect the evidence needed for the next milestone.")
    # Demo sequence
    add_rect(s, 0.74, 2.01, 5.56, 2.54, fill="card", line="mint", radius=True, line_width=0.8)
    pill(s, 0.98, 2.20, "DEMO SEQUENCE", color="mint", width=1.45, size=7.8)
    steps = [
        "Issue a sidecar for a known test audio/video file.",
        "Verify the original: signature status + exact hash match.",
        "Transcode a copy: show the hash mismatch without calling it fake.",
        "Show optional candidate score and a human-review/OCR sample.",
    ]
    for i, t in enumerate(steps):
        cy = 2.67 + i*0.40
        pill(s, 0.99, cy, f"0{i+1}", color="gold", width=0.42, size=7.5)
        text_box(s, 1.57, cy+0.015, 4.43, 0.38, t, size=9.9, color="paper", margin=0)

    add_rect(s, 6.53, 2.01, 6.07, 2.54, fill="card", line="blue", radius=True, line_width=0.8)
    pill(s, 6.77, 2.20, "VERIFIED PROJECT CHECKS", color="blue", width=2.00, size=7.8)
    text_box(s, 6.80, 2.67, 5.50, 0.46, "JavaScript / manifest / Python syntax checks · production build", size=11.4, color="paper", bold=True, margin=0)
    text_box(s, 6.80, 3.16, 5.50, 0.34, "7 / 7 credential tests passed", size=11.2, color="mint", bold=True, margin=0)
    text_box(s, 6.80, 3.53, 5.50, 0.34, "3 / 3 content-preflight tests passed", size=11.2, color="mint", bold=True, margin=0)
    text_box(s, 6.80, 3.91, 5.50, 0.41, "Not runtime-validated: Chrome extension, local backend, native Chromaprint, Tesseract OCR, and faster-whisper.", size=10.2, color="coral", margin=0)

    add_rect(s, 0.74, 4.80, 7.00, 1.36, fill="card2", line="gold", radius=True, line_width=0.75)
    pill(s, 0.98, 4.98, "INPUT NEEDED", color="gold", width=1.16, size=7.8)
    text_box(s, 0.99, 5.38, 6.48, 0.52, "Sanitized SBR API docs + test vectors; a Chrome test setup; and a decision on what—if anything—should be anchored on Monad.", size=11.2, color="paper", bold=True, margin=0)
    add_rect(s, 7.95, 4.80, 4.65, 1.36, fill="card", line="line", radius=True, line_width=0.65)
    text_box(s, 8.18, 4.99, 4.18, 0.78, "No entrant/winner claim. No Monad deployment or endorsement is represented. Content review is not legal or SEBI approval.", size=10.9, color="muted", margin=0)

    source_link(s, 0.78, 6.36, "Official brief: monad.xyz/developers/hackathons/metropolis", "https://monad.xyz/developers/hackathons/metropolis", width=5.6)
    source_link(s, 6.56, 6.36, "Official logo kit: monad.xyz/brand-and-media-kit", "https://monad.xyz/brand-and-media-kit", width=5.2)
    text_box(s, 0.80, 6.73, 11.4, 0.18, "Prepared 13 Oct 2026 · RelayProof prototype notes: LIVE_VERIFIER.md · Official Monad logo/Metropolis wordmark used for ecosystem identification only.", size=8.3, color="dim", font=MONO, margin=0)
    return s


def build() -> Path:
    ensure_assets()
    prs = Presentation()
    prs.slide_width = Inches(SW); prs.slide_height = Inches(SH)
    prs.core_properties.title = "RelayProof — Monad Metropolis Track 04 Problem-Fit Deck"
    prs.core_properties.subject = "Seven-slide prototype overview, official prompt fit, features, architecture, limitations and validation"
    prs.core_properties.author = "RelayProof prototype"
    prs.core_properties.keywords = "RelayProof, Monad, Metropolis, Track 04, media provenance"
    slide1(prs)
    slide2(prs)
    slide3(prs)
    slide4(prs)
    slide5(prs)
    slide6(prs)
    slide7(prs)
    prs.save(PPTX_PATH)
    print(f"Saved {PPTX_PATH} ({PPTX_PATH.stat().st_size:,} bytes, {len(prs.slides)} slides)")
    return PPTX_PATH


if __name__ == "__main__":
    build()
